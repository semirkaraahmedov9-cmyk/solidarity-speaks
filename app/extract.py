"""Извличане на текст от PDF, EPUB, DOCX и TXT файлове и разделяне на глави.

Редът на търсене на глави:
1. Структурата на самия документ (отметки в PDF, стилове „Heading“ в DOCX, съдържание в EPUB).
2. Редове като „Глава 1“, „ГЛАВА ПЪРВА“, „Chapter 1“, римски цифри.
3. Ако няма нищо: равни части от около 4500 думи, разделени между абзаците.
"""

import codecs
import io
import posixpath
import re
import urllib.parse
from dataclasses import dataclass, field
from pathlib import PurePath

import ebooklib
import pymupdf
from bs4 import BeautifulSoup
from docx import Document
from docx.table import Table
from ebooklib import epub

from . import config


class ExtractError(Exception):
    """Грешка, която показваме на потребителя. Кодът съответства на error_<код> в static/i18n/."""

    def __init__(self, code: str):
        super().__init__(code)
        self.code = code


@dataclass
class Chapter:
    # Заглавие от документа; None, ако няма (страницата показва „Част 3“)
    title: str | None
    paragraphs: list[str]
    # True за текста преди първата глава (заглавна страница, предговор и т.н.)
    intro: bool = False

    @property
    def words(self) -> int:
        return sum(len(paragraph.split()) for paragraph in self.paragraphs)

    @property
    def text(self) -> str:
        return "\n\n".join(self.paragraphs)


@dataclass
class Book:
    chapters: list[Chapter]
    # Как са намерени главите: "structure", "headings" или "parts"
    source: str


@dataclass
class Heading:
    index: int  # номер на абзаца, от който започва главата
    title: str
    level: int = 1  # 1 = основно заглавие, 2 = подзаглавие и т.н.


@dataclass
class RawText:
    """Резултатът от четенето на файл: абзаци + заглавия от структурата на документа."""

    paragraphs: list[str] = field(default_factory=list)
    headings: list[Heading] = field(default_factory=list)


# --- Основна функция ---


def extract_book(filename: str, data: bytes) -> Book:
    """Чете файла и го разделя на глави. При проблем хвърля ExtractError."""
    extension = PurePath(filename).suffix.lower()
    readers = {".pdf": _read_pdf, ".epub": _read_epub, ".docx": _read_docx, ".txt": _read_txt}
    if extension not in readers:
        raise ExtractError("unsupported_type")
    if not data.strip():
        raise ExtractError("empty_file")

    raw = readers[extension](data)
    paragraphs = raw.paragraphs
    total_words = sum(len(paragraph.split()) for paragraph in paragraphs)
    if total_words == 0:
        raise ExtractError("empty_file")

    # 1. Структурата на документа
    headings = _choose_headings(raw.headings, total_words)
    if len(headings) >= 2:
        return Book(_split_at_headings(paragraphs, headings), "structure")

    # 2. Редове, които приличат на заглавия на глави
    headings = [
        Heading(index, paragraph) for index, paragraph in enumerate(paragraphs) if is_chapter_heading(paragraph)
    ]
    if len(headings) >= 2:
        return Book(_split_at_headings(paragraphs, headings), "headings")

    # 3. Равни части
    return Book(_split_into_parts(paragraphs), "parts")


# --- Помощни функции за текст ---


def _clean(text: str) -> str:
    """Премахва излишните интервали и нови редове."""
    return " ".join(text.split())


def _join_lines(lines: list[str]) -> str:
    """Слепва редовете на един абзац. Дума, пренесена с тире на нов ред, се слепва обратно."""
    result = ""
    for line in lines:
        line = line.strip()
        if not line:
            continue
        if result.endswith("-") and len(result) > 1 and result[-2].isalpha() and line[0].islower():
            result = result[:-1] + line
        elif result:
            result += " " + line
        else:
            result = line
    return _clean(result)


def _normalize(text: str) -> str:
    """Само малки букви и цифри — за сравняване на заглавия."""
    return "".join(char for char in text.lower() if char.isalnum())


# --- Разпознаване на заглавия по текста ---

# Поредни числа с думи: „първа“, „втора“, …, „двадесет и първа“
_BG_ORDINALS = (
    r"(?:първ|втор|трет|четвърт|пет|шест|седм|осм|девет|десет|единадесет|единайсет|"
    r"дванадесет|дванайсет|тринадесет|тринайсет|четиринадесет|четиринайсет|петнадесет|"
    r"петнайсет|шестнадесет|шестнайсет|седемнадесет|седемнайсет|осемнадесет|осемнайсет|"
    r"деветнадесет|деветнайсет|двадесет|двайсет|тридесет|трийсет)\w*(?:\s+и\s+\w+)?"
)
_EN_NUMBERS = (
    r"(?:one|two|three|four|five|six|seven|eight|nine|ten|eleven|twelve|thirteen|fourteen|"
    r"fifteen|sixteen|seventeen|eighteen|nineteen|twenty|thirty|forty|fifty)(?:[\s-]\w+)?"
)
# Римски числа от I до MMMCMXCIX
_ROMAN = r"(?=[IVXLCDM])M{0,3}(?:CM|CD|D?C{0,3})(?:XC|XL|L?X{0,3})(?:IX|IV|V?I{0,3})"

# „Глава 1“, „ГЛАВА ПЪРВА“, „Chapter 12“, „Част втора“, „Chapter IV. Заглавие“…
_NUMBERED_HEADING = re.compile(
    rf"^(?:глава|част|chapter|part)\s+(?:\d+|{_ROMAN}|{_BG_ORDINALS}|{_EN_NUMBERS})\b",
    re.IGNORECASE,
)
# Ред само с римско число: „IV“ или „IV.“ (само главни букви, за да не хващаме думи)
_ROMAN_HEADING = re.compile(rf"^{_ROMAN}\.?$")
# Пролог, епилог и подобни, сами на реда
_NAMED_HEADING = re.compile(
    r"^(?:пролог|епилог|предговор|послеслов|увод|prologue|epilogue|preface|introduction|afterword)\.?$",
    re.IGNORECASE,
)

# Заглавията са кратки; по-дълъг ред е обикновено изречение
_MAX_HEADING_CHARS = 100


def is_chapter_heading(line: str) -> bool:
    """Проверява дали ред прилича на заглавие на глава."""
    line = line.strip()
    if not line or len(line) > _MAX_HEADING_CHARS:
        return False
    return bool(_NUMBERED_HEADING.match(line) or _ROMAN_HEADING.match(line) or _NAMED_HEADING.match(line))


# --- Разделяне на глави ---


def _choose_headings(headings: list[Heading], total_words: int) -> list[Heading]:
    """Избира колко нива на заглавията да използваме.

    Започваме от основните заглавия. Ако те са по-малко от 2 или главите са твърде дълги
    (напр. „Част първа“, „Част втора“), добавяме и следващото ниво (главите вътре в частите).
    """
    levels = sorted({heading.level for heading in headings})
    chosen: list[Heading] = []
    for level in levels:
        chosen = [heading for heading in headings if heading.level <= level]
        if len(chosen) >= 2 and total_words / len(chosen) <= config.MAX_AVERAGE_CHAPTER_WORDS:
            break
    return chosen


def _split_at_headings(paragraphs: list[str], headings: list[Heading]) -> list[Chapter]:
    """Разделя абзаците на глави там, където започват заглавията."""
    # Подреждаме по позиция; две заглавия на едно място — пазим първото
    starts: list[Heading] = []
    for heading in sorted(headings, key=lambda h: h.index):
        if not starts or heading.index > starts[-1].index:
            starts.append(heading)

    chapters = []
    # Текстът преди първата глава не се губи: става „Начало“
    if starts[0].index > 0:
        chapters.append(Chapter(None, paragraphs[: starts[0].index], intro=True))
    for number, heading in enumerate(starts):
        end = starts[number + 1].index if number + 1 < len(starts) else len(paragraphs)
        chapters.append(Chapter(_clean(heading.title) or None, paragraphs[heading.index : end]))

    # Глава, която е само заглавие (напр. „Част първа“ точно преди „Глава 1“),
    # се слива със следващата, за да не стои празна в списъка
    merged: list[Chapter] = []
    carry: list[str] = []
    for position, chapter in enumerate(chapters):
        chapter.paragraphs = carry + chapter.paragraphs
        carry = []
        is_last = position == len(chapters) - 1
        if not chapter.intro and not is_last and len(chapter.paragraphs) == 1 and chapter.words <= 15:
            carry = chapter.paragraphs
            continue
        merged.append(chapter)
    return [chapter for chapter in merged if chapter.words > 0]


def _split_into_parts(paragraphs: list[str]) -> list[Chapter]:
    """Разделя текста на почти равни части от около PART_TARGET_WORDS думи, между абзаците."""
    # Много дълъг абзац (напр. текст без нови редове) се разделя на изречения
    pieces: list[str] = []
    for paragraph in paragraphs:
        if len(paragraph.split()) > 1000:
            pieces.extend(sentence for sentence in re.split(r"(?<=[.!?…])\s+", paragraph) if sentence)
        else:
            pieces.append(paragraph)

    total = sum(len(piece.split()) for piece in pieces)
    count = max(1, round(total / config.PART_TARGET_WORDS))
    target = total / count

    parts: list[list[str]] = [[]]
    done = 0
    for piece in pieces:
        if parts[-1] and len(parts) < count and done >= target * len(parts):
            parts.append([])
        parts[-1].append(piece)
        done += len(piece.split())
    return [Chapter(None, part) for part in parts]


# --- TXT ---


def _decode_text(data: bytes) -> str:
    """Разпознава кодирането: UTF-8 (най-често) или Windows-1251 (стари български файлове)."""
    if data.startswith(codecs.BOM_UTF8):
        return data[len(codecs.BOM_UTF8) :].decode("utf-8", errors="replace")
    if data.startswith((codecs.BOM_UTF16_LE, codecs.BOM_UTF16_BE)):
        return data.decode("utf-16", errors="replace")
    try:
        return data.decode("utf-8")
    except UnicodeDecodeError:
        return data.decode("cp1251", errors="replace")


def _read_txt(data: bytes) -> RawText:
    text = _decode_text(data)
    lines = text.splitlines()
    # Ако има празни редове между абзаците, те разделят абзаците
    # (редовете вътре в абзаца се слепват). Иначе всеки ред е отделен абзац.
    blank_separated = re.search(r"\S[ \t]*\r?\n[ \t]*\r?\n", text) is not None

    raw = RawText()
    current: list[str] = []

    def finish_paragraph():
        if current:
            raw.paragraphs.append(_join_lines(current))
            current.clear()

    for line in lines:
        stripped = line.strip()
        if not stripped:
            finish_paragraph()
        elif not blank_separated or is_chapter_heading(stripped):
            # Заглавието винаги е отделен абзац, дори без празен ред след него
            finish_paragraph()
            raw.paragraphs.append(_clean(stripped))
        else:
            current.append(stripped)
    finish_paragraph()
    return raw


# --- DOCX ---


def _docx_heading_level(style_name: str) -> int | None:
    """„Heading 2“ → 2, „Title“ → 0 (заглавие на документа), иначе None."""
    name = style_name.strip().lower()
    if name == "title":
        return 0
    match = re.match(r"(?:heading|заглавие)\s*(\d)", name)
    return int(match.group(1)) if match else None


def _read_docx(data: bytes) -> RawText:
    try:
        document = Document(io.BytesIO(data))
    except Exception as error:
        raise ExtractError("unreadable_file") from error

    raw = RawText()
    # Абзаците и таблиците, в реда, в който са в документа
    for block in document.iter_inner_content():
        if isinstance(block, Table):
            for row in block.rows:
                previous_cell = None
                for cell in row.cells:
                    # Слетите клетки се повтарят; четем ги веднъж
                    if previous_cell is not None and cell._tc is previous_cell._tc:
                        continue
                    previous_cell = cell
                    text = _clean(cell.text)
                    if text:
                        raw.paragraphs.append(text)
            continue

        text = _clean(block.text)
        if not text:
            continue
        style_name = block.style.name if block.style is not None else ""
        level = _docx_heading_level(style_name or "")
        if level is not None:
            raw.headings.append(Heading(len(raw.paragraphs), text, level))
        raw.paragraphs.append(text)
    return raw


# --- EPUB ---

# HTML елементи, които съдържат абзаци текст
_BLOCK_TAGS = [
    "p", "h1", "h2", "h3", "h4", "h5", "h6", "li", "blockquote", "pre",
    "div", "dd", "dt", "td", "th", "figcaption", "caption",
]  # fmt: skip
_BLOCK_SET = set(_BLOCK_TAGS)


def _flatten_toc(toc, level: int = 1):
    """Превръща вложеното съдържание на EPUB в плосък списък (ниво, заглавие, адрес)."""
    for entry in toc:
        if isinstance(entry, tuple):
            section, children = entry
            if getattr(section, "href", None):
                yield level, section.title, section.href
            yield from _flatten_toc(children, level + 1)
        elif getattr(entry, "href", None):
            yield level, entry.title, entry.href


def _read_epub(data: bytes) -> RawText:
    try:
        book = epub.read_epub(io.BytesIO(data), {"ignore_ncx": False})
    except Exception as error:
        raise ExtractError("unreadable_file") from error

    raw = RawText()
    # Къде започва всеки файл и всяка котва (#id) в списъка с абзаци
    positions: dict[tuple[str, str], int] = {}
    file_names: list[str] = []
    # Първото заглавие (h1–h3) във всеки файл — ако книгата няма съдържание
    spine_headings: list[Heading] = []

    # Файловете в реда на четене (spine)
    for item_id, _linear in book.spine:
        item = book.get_item_with_id(item_id)
        # Пропускаме страницата със съдържанието (nav), за да не се чете на глас
        if item is None or item.get_type() != ebooklib.ITEM_DOCUMENT or not item.is_chapter():
            continue
        name = item.get_name()
        file_names.append(name)
        positions[(name, "")] = len(raw.paragraphs)
        file_start = len(raw.paragraphs)
        first_heading = None

        soup = BeautifulSoup(item.get_content(), "html.parser")
        for tag in soup(["script", "style"]):
            tag.decompose()
        for line_break in soup.find_all("br"):
            line_break.replace_with(" ")

        body = soup.body or soup
        for element in body.find_all(True):
            if element.get("id"):
                positions.setdefault((name, element["id"]), len(raw.paragraphs))
            # Вземаме само най-вътрешните блокове, за да не четем текста два пъти
            if element.name in _BLOCK_SET and element.find(_BLOCK_TAGS) is None:
                text = _clean(element.get_text())
                if not text:
                    continue
                if first_heading is None and element.name in ("h1", "h2", "h3"):
                    first_heading = text
                raw.paragraphs.append(text)

        if len(raw.paragraphs) > file_start:
            spine_headings.append(Heading(file_start, first_heading or ""))

    # Свързваме съдържанието (TOC) с позициите в текста
    for level, title, href in _flatten_toc(book.toc):
        file_part, _, anchor = urllib.parse.unquote(href).partition("#")
        base_name = posixpath.basename(file_part)
        name = next((n for n in file_names if posixpath.basename(n) == base_name), None)
        if name is None:
            continue
        index = positions.get((name, anchor), positions[(name, "")])
        if index < len(raw.paragraphs):
            raw.headings.append(Heading(index, title or "", level))

    # Без използваемо съдържание: всеки файл от книгата е отделна глава
    if len({heading.index for heading in raw.headings}) < 2:
        raw.headings = spine_headings
    return raw


# --- PDF ---


def _find_title(paragraphs: list[str], start: int, end: int, title: str) -> int:
    """Търси на страницата абзаца, с който започва главата (отметката сочи само страницата)."""
    key = _normalize(title)
    if key:
        for index in range(start, end):
            text = _normalize(paragraphs[index])
            if text.startswith(key[:40]) or (len(text) >= 3 and key.startswith(text)):
                return index
    return start


def _read_pdf(data: bytes) -> RawText:
    try:
        document = pymupdf.open(stream=data, filetype="pdf")
    except Exception as error:
        raise ExtractError("unreadable_file") from error
    if document.needs_pass:
        raise ExtractError("pdf_password")

    raw = RawText()
    page_starts: list[int] = []  # първият абзац на всяка страница
    total_chars = 0
    for page in document:
        page_starts.append(len(raw.paragraphs))
        # Блоковете текст на страницата обикновено съответстват на абзаци
        for block in page.get_text("blocks"):
            if block[6] != 0:  # 1 = картинка
                continue
            text = _join_lines(block[4].splitlines())
            if text:
                raw.paragraphs.append(text)
                total_chars += len(text)
    page_starts.append(len(raw.paragraphs))

    # Почти без текст: сканиран документ (картинки) или празен файл
    if total_chars < config.SCANNED_PDF_MIN_CHARS_PER_PAGE * max(1, document.page_count):
        has_images = any(page.get_images() for page in document)
        raise ExtractError("scanned_pdf" if has_images else "empty_file")

    # Отметките (bookmarks): [ниво, заглавие, страница]
    for level, title, page_number in document.get_toc(simple=True):
        if not 1 <= page_number <= document.page_count:
            continue
        start, end = page_starts[page_number - 1], page_starts[page_number]
        raw.headings.append(Heading(_find_title(raw.paragraphs, start, end, title), title, level))
    return raw
