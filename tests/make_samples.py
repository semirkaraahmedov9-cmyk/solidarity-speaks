"""Създава малките примерни файлове в tests/samples/, които тестовете използват.

Файловете вече са в git; този скрипт е нужен само ако искате да ги създадете отново:
    .venv\\Scripts\\python.exe tests\\make_samples.py
"""

from pathlib import Path

import pymupdf
from docx import Document
from ebooklib import epub

SAMPLES = Path(__file__).resolve().parent / "samples"

TITLE = "Малка книга за проба"
CHAPTERS = [
    (
        "Глава първа",
        [
            "Беше хубав пролетен ден. Децата играеха в двора на училището.",
            "Учителката ги гледаше от прозореца и се усмихваше.",
        ],
    ),
    (
        "Глава втора",
        [
            "Звънецът би и децата се затичаха към класната стая.",
            "Урокът по литература започваше с приказка за смелото момиче.",
        ],
    ),
    (
        "Глава трета",
        [
            "Когато приказката свърши, едно момче вдигна ръка.",
            "Смелостта не е липса на страх, а решението да направиш правилното нещо.",
        ],
    ),
]


def make_txt():
    lines = [TITLE, ""]
    for title, paragraphs in CHAPTERS:
        # Заглавие без празен ред след него — както в много истински TXT файлове
        lines.append(title.upper())
        for paragraph in paragraphs:
            lines += [paragraph, ""]
    text = "\n".join(lines)
    (SAMPLES / "sample.txt").write_text(text, encoding="utf-8")
    # Същият текст в старото кодиране Windows-1251
    (SAMPLES / "sample_cp1251.txt").write_bytes(text.replace("\n", "\r\n").encode("cp1251"))


def make_docx():
    document = Document()
    document.add_heading(TITLE, level=0)
    for number, (title, paragraphs) in enumerate(CHAPTERS):
        document.add_heading(title, level=1)
        for paragraph in paragraphs:
            document.add_paragraph(paragraph)
        if number == 1:
            # Таблица във втората глава — текстът ѝ също трябва да се прочете
            table = document.add_table(rows=1, cols=2)
            table.cell(0, 0).text = "Ученик"
            table.cell(0, 1).text = "Оценка"
    document.save(SAMPLES / "sample.docx")


def make_epub():
    book = epub.EpubBook()
    book.set_identifier("solidarity-speaks-sample")
    book.set_title(TITLE)
    book.set_language("bg")
    items = []
    for number, (title, paragraphs) in enumerate(CHAPTERS, start=1):
        item = epub.EpubHtml(title=title, file_name=f"chapter{number}.xhtml", lang="bg")
        body = "".join(f"<p>{paragraph}</p>" for paragraph in paragraphs)
        item.content = f"<html><body><h1>{title}</h1>{body}</body></html>"
        book.add_item(item)
        items.append(item)
    book.toc = [epub.Link(item.file_name, item.title, f"ch{n}") for n, item in enumerate(items, start=1)]
    book.add_item(epub.EpubNcx())
    book.add_item(epub.EpubNav())
    book.spine = ["nav", *items]
    epub.write_epub(SAMPLES / "sample.epub", book)


def _pdf_page(document, html):
    page = document.new_page()
    page.insert_htmlbox(pymupdf.Rect(50, 50, 545, 790), html)


def _chapter_html(title, paragraphs):
    return f"<h2>{title}</h2>" + "".join(f"<p>{paragraph}</p>" for paragraph in paragraphs)


def _save(document, path):
    # Записваме само използваните букви от шрифта, за да е файлът малък
    document.subset_fonts()
    document.save(path, garbage=4, deflate=True)


def make_pdfs():
    # PDF с отметки (bookmarks). Глави 2 и 3 са на една и съща страница.
    document = pymupdf.open()
    _pdf_page(document, f"<h1>{TITLE}</h1>" + _chapter_html(*CHAPTERS[0]))
    _pdf_page(document, _chapter_html(*CHAPTERS[1]) + _chapter_html(*CHAPTERS[2]))
    document.set_toc([[1, CHAPTERS[0][0], 1], [1, CHAPTERS[1][0], 2], [1, CHAPTERS[2][0], 2]])
    _save(document, SAMPLES / "sample_bookmarks.pdf")

    # PDF без отметки: главите се разпознават по текста „ГЛАВА ПЪРВА“ и т.н.
    document = pymupdf.open()
    _pdf_page(document, f"<h1>{TITLE}</h1>")
    for title, paragraphs in CHAPTERS:
        _pdf_page(document, _chapter_html(title.upper(), paragraphs))
    _save(document, SAMPLES / "sample_headings.pdf")

    # „Сканиран“ PDF: само картинка, без текст
    document = pymupdf.open()
    page = document.new_page()
    picture = pymupdf.Pixmap(pymupdf.csRGB, pymupdf.IRect(0, 0, 200, 300), False)
    picture.clear_with(220)
    page.insert_image(pymupdf.Rect(50, 50, 545, 790), pixmap=picture)
    _save(document, SAMPLES / "sample_scanned.pdf")


if __name__ == "__main__":
    SAMPLES.mkdir(exist_ok=True)
    make_txt()
    make_docx()
    make_epub()
    make_pdfs()
    print("Примерните файлове са създадени в", SAMPLES)
